#!/usr/bin/env python3
"""Le resultados.json (+ correcoes.json opcional da busca web) e escreve a aba
consolidada 'Validacao' no xlsx, com coluna de Lote. Preserva as abas originais."""
import json, os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

XLSX = "base_contatos.xlsx"
OUT = "base_contatos_validada.xlsx"
BATCH = 100

HDR = ["Lote", "Dominio_base", "Qtd_Contatos", "Empresa_base", "DNS_ativo",
       "HTTP_status", "Site_responde", "Funciona", "Dominio_corrigido",
       "URL_final", "Metodo_correcao", "Observacao_validacao"]

def load(path, default=None):
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    return default

def main():
    results = load("resultados.json")
    correcoes = load("correcoes.json", {})  # {host: {dominio_corrigido, funciona, metodo, obs, http, responde, final_url}}

    # aplica correcoes da busca web sobre os resultados
    for r in results:
        c = correcoes.get(r["host"])
        if not c:
            continue
        if c.get("dominio_corrigido"):
            r["dominio_corrigido"] = c["dominio_corrigido"]
            r["funciona"] = c.get("funciona", "SIM (corrigido)")
            r["metodo_correcao"] = c.get("metodo", "busca web")
            r["http"] = c.get("http", r["http"])
            r["responde"] = c.get("responde", r["responde"])
            r["final_url"] = c.get("final_url", r["final_url"])
            r["obs"] = c.get("obs", r["obs"])
        else:
            # busca web confirmou que nao existe site valido
            r["funciona"] = "NAO"
            r["metodo_correcao"] = "busca web"
            if c.get("obs"):
                r["obs"] = c["obs"]

    wb = openpyxl.load_workbook(XLSX)
    if "Validacao" in wb.sheetnames:
        del wb["Validacao"]
    ws = wb.create_sheet("Validacao")

    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    fill_ok = PatternFill("solid", fgColor="C6EFCE")
    fill_fix = PatternFill("solid", fgColor="FFEB9C")
    fill_bad = PatternFill("solid", fgColor="FFC7CE")

    ws.append(HDR)
    for c in range(1, len(HDR) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill; cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for idx, r in enumerate(results):
        if r is None:
            continue
        lote = idx // BATCH + 1
        row = [lote, r["dominio"], r["qtd"], r["empresa"], r["dns"],
               r["http"], r["responde"], r["funciona"], r["dominio_corrigido"],
               r["final_url"], r["metodo_correcao"], r["obs"]]
        ws.append(row)
        rownum = ws.max_row
        f = r["funciona"]
        fill = fill_ok if f == "SIM" else fill_fix if f.startswith("SIM (") else fill_bad
        ws.cell(row=rownum, column=8).fill = fill

    widths = [6, 30, 12, 32, 11, 11, 18, 16, 30, 38, 16, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    wb.save(OUT)
    from collections import Counter
    validos = [r for r in results if r is not None]
    print(f"Salvo: {OUT}")
    print("Total linhas:", len(validos))
    print("funciona:", Counter(r["funciona"] for r in validos))
    print("Lotes:", (len(results) - 1) // BATCH + 1)

if __name__ == "__main__":
    main()
