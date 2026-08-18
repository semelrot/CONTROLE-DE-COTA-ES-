#!/usr/bin/env python3
"""Consolida resultados.json + correcoes.json + pesquisa/part_*.json +
research_queue.json numa unica planilha ("dados"), uma linha por dominio.

Reproduz a mesma logica de merge de correcoes que escrever.py aplica sobre a
aba 'Validacao', mas sem depender de base_contatos.xlsx (nao versionado) --
os campos Tipo/Subtipo/Consome_aco/Obs_base vem de research_queue.json, que ja
tinha sido cruzado com essa planilha numa sessao anterior (build_queue.py).

Saida: controle_cotacoes_dados.xlsx, aba unica "dados", como Tabela do Excel.
"""
import glob
import json
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

SAIDA = "controle_cotacoes_dados.xlsx"

COLUNAS = [
    ("lote",              "Lote"),
    ("idx",               "Índice"),
    ("dominio",           "Domínio (base)"),
    ("empresa",           "Empresa"),
    ("qtd",               "Qtd. Contatos"),
    ("dns",               "DNS ativo"),
    ("http",              "Status HTTP"),
    ("responde",          "Site responde"),
    ("funciona",          "Funciona"),
    ("dominio_corrigido", "Domínio corrigido"),
    ("final_url",         "URL final"),
    ("metodo_correcao",   "Método da correção"),
    ("obs",               "Observação (validação)"),
    ("tipo_base",         "Tipo (base)"),
    ("subtipo_base",      "Subtipo (base)"),
    ("consome_base",      "Consome aço (base)"),
    ("obs_base",          "Observação (base)"),
    ("segmento",          "Segmento (pesquisa)"),
    ("principal_produto", "Principal produto (pesquisa)"),
    ("fabrica_brasil",    "Fábrica no Brasil (pesquisa)"),
    ("usa_aco_carbono",   "Usa aço carbono (pesquisa)"),
    ("justificativa_aco", "Justificativa aço (pesquisa)"),
    ("fonte",             "Fonte da pesquisa"),
]

FONTE = "Arial"
BATCH = 100


def carregar(caminho, padrao=None):
    try:
        with open(caminho, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return padrao


def carregar_pesquisa():
    """Junta pesquisa/part_*.json num dict idx -> registro."""
    out = {}
    for fp in sorted(glob.glob("pesquisa/part_*.json")):
        for e in json.load(open(fp, encoding="utf-8")):
            if e.get("idx") is not None:
                out[e["idx"]] = e
    return out


def aplicar_correcoes(results, correcoes):
    """Mesma logica de escrever.py: correcoes.json sobrepoe resultados.json."""
    for r in results:
        c = correcoes.get(r["host"])
        if not c:
            continue
        if c.get("dominio_corrigido"):
            r["dominio_corrigido"] = c["dominio_corrigido"]
            r["funciona"] = c.get("funciona", "SIM (corrigido)")
            r["metodo_correcao"] = c.get("metodo", "busca web")
            r["http"] = c.get("http", r["http"])
            r["responde"] = c.get("responde", r["responde"])
            r["final_url"] = c.get("final_url", r["final_url"])
            r["obs"] = c.get("obs", r["obs"])
        else:
            # busca web confirmou que nao existe site valido
            r["funciona"] = "NAO"
            r["metodo_correcao"] = "busca web"
            if c.get("obs"):
                r["obs"] = c["obs"]
    return results


def montar_linhas():
    results = carregar("resultados.json", [])
    correcoes = carregar("correcoes.json", {})
    pesquisa = carregar_pesquisa()
    queue_por_idx = {q["idx"]: q for q in carregar("research_queue.json", [])}

    results = aplicar_correcoes(results, correcoes)

    linhas = []
    for idx, r in enumerate(results):
        if r is None:
            continue
        p = pesquisa.get(idx, {})
        q = queue_por_idx.get(idx, {})
        linhas.append({
            "lote": idx // BATCH + 1,
            "idx": idx,
            "dominio": r.get("dominio"),
            "empresa": r.get("empresa"),
            "qtd": r.get("qtd"),
            "dns": r.get("dns"),
            "http": r.get("http"),
            "responde": r.get("responde"),
            "funciona": r.get("funciona"),
            "dominio_corrigido": r.get("dominio_corrigido"),
            "final_url": r.get("final_url"),
            "metodo_correcao": r.get("metodo_correcao"),
            "obs": r.get("obs"),
            "tipo_base": q.get("tipo_base"),
            "subtipo_base": q.get("subtipo_base"),
            "consome_base": q.get("consome_base"),
            "obs_base": q.get("obs_base"),
            "segmento": p.get("segmento"),
            "principal_produto": p.get("principal_produto"),
            "fabrica_brasil": p.get("fabrica_brasil"),
            "usa_aco_carbono": p.get("usa_aco_carbono"),
            "justificativa_aco": p.get("justificativa_aco"),
            "fonte": p.get("fonte"),
        })
    return linhas


def escrever_planilha(linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "dados"

    chaves = [c[0] for c in COLUNAS]
    cabecalhos = [c[1] for c in COLUNAS]
    ws.append(cabecalhos)
    for cel in ws[1]:
        cel.font = Font(name=FONTE, bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="1F4E78")
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    for linha in linhas:
        ws.append([linha.get(k) for k in chaves])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cel in row:
            cel.font = Font(name=FONTE)
            cel.alignment = Alignment(vertical="center")

    larguras = {
        "Lote": 7, "Índice": 8, "Domínio (base)": 28, "Empresa": 32,
        "Qtd. Contatos": 12, "DNS ativo": 11, "Status HTTP": 11,
        "Site responde": 20, "Funciona": 16, "Domínio corrigido": 28,
        "URL final": 38, "Método da correção": 16, "Observação (validação)": 48,
        "Tipo (base)": 18, "Subtipo (base)": 16, "Consome aço (base)": 16,
        "Observação (base)": 30, "Segmento (pesquisa)": 20,
        "Principal produto (pesquisa)": 34, "Fábrica no Brasil (pesquisa)": 30,
        "Usa aço carbono (pesquisa)": 16, "Justificativa aço (pesquisa)": 46,
        "Fonte da pesquisa": 26,
    }
    for i, nome in enumerate(cabecalhos, 1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(nome, 18)

    ultima_col = get_column_letter(len(cabecalhos))
    ref = f"A1:{ultima_col}{ws.max_row}"
    tabela = Table(displayName="Dados", ref=ref)
    tabela.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(tabela)

    wb.save(SAIDA)


def main():
    linhas = montar_linhas()
    escrever_planilha(linhas)

    print(f"Salvo: {SAIDA}")
    print(f"Total de linhas: {len(linhas)}")
    print("Funciona:", dict(Counter(l["funciona"] for l in linhas)))
    print("Com Tipo (base) preenchido:", sum(1 for l in linhas if l["tipo_base"]))
    print("Com pesquisa (segmento) preenchido:", sum(1 for l in linhas if l["segmento"]))


if __name__ == "__main__":
    main()
